"""
Pagina de gestion de promociones mensuales desde Excel.
"""

import streamlit as st
import openpyxl
import io
from storage import (
    get_proyectos,
    get_promociones,
    add_promocion,
    get_promociones_by_mes,
    get_meses_promociones,
    delete_promocion,
)
from utils import analizar_excel_promociones, verificar_api_key

st.markdown("# 🎯 Promociones del Mes")
st.markdown("Sube un archivo Excel con las promociones mensuales y el AI analizara los proyectos para registrar los beneficios.")

api_ok = verificar_api_key()
if not api_ok:
    st.error("⚠️ Configura tu API Key en el archivo .env para el proveedor seleccionado.")
    st.stop()

proyectos = get_proyectos()
if not proyectos:
    st.warning("📋 No hay proyectos registrados. Ve a la seccion **Proyectos** para agregar uno primero.")
    st.stop()

# ===== FORMULARIO DE CARGA =====
st.markdown("### 📥 Cargar promociones del mes")

col1, col2 = st.columns([1, 2])
with col1:
    nombre_mes = st.text_input(
        "Nombre de la promocion",
        placeholder="Ej: Promocion Julio 2026, Mes del Nino...",
        help="Ponle un nombre descriptivo a esta promocion.",
    )
with col2:
    archivo_excel = st.file_uploader(
        "Selecciona el archivo Excel (.xlsx)",
        type=["xlsx", "xls"],
        help="Sube un Excel con las promociones del mes. El AI lo analizara automaticamente.",
    )

if archivo_excel and nombre_mes:
    with st.status("🔍 Analizando archivo...", expanded=True) as status:
        st.write("📖 Leyendo contenido del Excel...")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(archivo_excel.read()), data_only=True)
            hojas_texto = []

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                st.write(f"  → Hoja: **{sheet_name}** ({ws.max_row} filas × {ws.max_column} columnas)")

                # Obtener encabezados
                rows_iter = ws.iter_rows(values_only=True)
                headers = next(rows_iter, None)
                
                filas = []
                if headers:
                    header_str = " | ".join(str(h) if h is not None else "" for h in headers)
                    filas.append(f"ENCABEZADOS: {header_str}")
                
                for row in rows_iter:
                    if any(cell is not None for cell in row):
                        # Formatear cada fila de forma que sea clara para el AI
                        fila_detalles = []
                        for i, cell in enumerate(row):
                            header_name = str(headers[i]) if headers and i < len(headers) else f"Col{i+1}"
                            val = str(cell) if cell is not None else ""
                            fila_detalles.append(f"{header_name}: {val}")
                        filas.append(" - ".join(fila_detalles))

                if filas:
                    hojas_texto.append(f"=== Hoja: {sheet_name} ===\n" + "\n".join(filas))

            texto_completo = "\n\n".join(hojas_texto)

            st.write(f"📄 Total caracteres extraidos: {len(texto_completo):,}")
            
            with st.expander("🔍 Ver texto extraido del Excel (Debug)", expanded=False):
                st.code(texto_completo)

            if len(texto_completo) < 20:
                st.error("El archivo parece estar vacio o no tiene datos legibles.")
                st.stop()

        except Exception as e:
            st.error(f"Error al leer el Excel: {str(e)}")
            st.stop()

        st.write("🤖 Enviando a Nemotron Nano Omni para analisis...")
        resultados = analizar_excel_promociones(texto_completo, nombre_mes)

        if resultados is None:
            st.error("Error de conexion con el proveedor AI. Verifica tu API Key.")
            st.stop()

        if not resultados:
            st.warning("⚠️ El AI no pudo extraer proyectos en formato JSON. Intenta subir el archivo de nuevo o verifica el formato del Excel.")
            st.info("💡 Consejo: Asegúrate de que los nombres de los proyectos (como 'Rosas 1444') estén en celdas claras.")
            st.stop()

        proyectos_dict = {p["nombre"].lower().strip(): p for p in proyectos}

        status.update(label="✅ Analisis completado", state="complete")

    # ===== MOSTRAR RESULTADOS =====
    st.markdown("### 📊 Resultados del analisis")

    proyectos_encontrados = 0
    proyectos_sin_match = 0

    for item in resultados:
        nombre_proyecto = item.get("nombre_proyecto", "").strip()
        desc_promocion = item.get("promocion", "").strip()

        if not nombre_proyecto:
            continue

        nombre_lower = nombre_proyecto.lower().strip()
        proyecto_match = proyectos_dict.get(nombre_lower)

        if not proyecto_match:
            for p_nombre, p in proyectos_dict.items():
                if nombre_lower in p_nombre or p_nombre in nombre_lower:
                    proyecto_match = p
                    break

        if proyecto_match:
            proyectos_encontrados += 1
            col1, col2 = st.columns([1, 3])
            with col1:
                st.success(f"✅ **{proyecto_match['nombre']}**")
            with col2:
                st.markdown(f"📌 **Promocion:** {desc_promocion}")
        else:
            proyectos_sin_match += 1
            col1, col2 = st.columns([1, 3])
            with col1:
                st.warning(f"⚠️ **{nombre_proyecto}**")
            with col2:
                st.markdown(f"📌 **Promocion:** {desc_promocion}")
                st.caption("Este proyecto no esta registrado en el sistema.")

    st.markdown("---")

    total_match = proyectos_encontrados
    if total_match > 0:
        st.markdown(f"**Se encontraron {total_match} proyecto(s) con match en el sistema.**")

        existentes = get_promociones_by_mes(nombre_mes)
        if existentes:
            st.warning(f"⚠️ Ya existen {len(existentes)} promociones registradas para **{nombre_mes}**. Al guardar, se agregaran las nuevas.")

        if st.button("💾 Guardar promociones", type="primary", use_container_width=True):
            guardadas = 0
            for item in resultados:
                nombre_proyecto = item.get("nombre_proyecto", "").strip()
                desc_promocion = item.get("promocion", "").strip()
                if not nombre_proyecto:
                    continue

                nombre_lower = nombre_proyecto.lower().strip()
                proyecto_match = proyectos_dict.get(nombre_lower)

                if not proyecto_match:
                    for p_nombre, p in proyectos_dict.items():
                        if nombre_lower in p_nombre or p_nombre in nombre_lower:
                            proyecto_match = p
                            break

                if proyecto_match and desc_promocion:
                    add_promocion(
                        proyecto_id=proyecto_match["id"],
                        nombre_proyecto=proyecto_match["nombre"],
                        mes=nombre_mes,
                        descripcion_promocion=desc_promocion,
                        archivo_original=archivo_excel.name,
                    )
                    guardadas += 1

            if guardadas > 0:
                st.success(f"✅ **{guardadas} promociones guardadas exitosamente para {nombre_mes}!**")
                st.balloons()
            else:
                st.warning("No se guardaron promociones. Verifica que los proyectos tengan match en el sistema.")

# ===== HISTORIAL DE PROMOCIONES =====
st.markdown("---")
st.markdown("### 📜 Historial de promociones cargadas")

meses = get_meses_promociones()
if meses:
    mes_seleccionado = st.selectbox("Filtrar por mes", ["Todos"] + meses)

    promociones = get_promociones()
    if mes_seleccionado != "Todos":
        promociones = [p for p in promociones if p["mes"] == mes_seleccionado]

    if promociones:
        for promo in promociones:
            with st.container(border=True):
                col1, col2, col3 = st.columns([2, 4, 1])
                with col1:
                    st.markdown(f"**{promo['nombre_proyecto']}**")
                    st.caption(f"📅 {promo['mes']}")
                with col2:
                    st.markdown(promo['descripcion_promocion'])
                    if promo.get('archivo_original'):
                        st.caption(f"📎 {promo['archivo_original']}")
                with col3:
                    if st.button("🗑️", key=f"del_{promo['id']}", help="Eliminar promocion"):
                        delete_promocion(promo['id'])
                        st.rerun()
    else:
        st.info("No hay promociones para el filtro seleccionado.")
else:
    st.info("📭 No hay promociones registradas aun. Sube un Excel mas arriba.")
